"""End-to-end tests for the QCAA word count pipeline.

Each test builds a .docx on the fly, runs the pipeline with
auto_accept=True (no interactive prompts), and asserts a known count.

These are deliberately hand-checkable: each fixture has a comment
explaining the expected count.
"""
from __future__ import annotations

import os
import tempfile
import unittest

from docx import Document as DocxDocument
from docx.shared import Pt


def _make_docx(paragraphs: list[str], path: str) -> None:
    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)


def _run(path: str) -> dict:
    """Run the pipeline end-to-end with auto-accept and return the result."""
    from qcaa_word_count.captions import detect_captions
    from qcaa_word_count.citations import (
        apply_citation_exclusions, collect_citations,
    )
    from qcaa_word_count.cli import _apply_span_decisions
    from qcaa_word_count.count import count_words
    from qcaa_word_count.footnotes import (
        apply_footnote_suggestions, read_footnotes_from_docx,
    )
    from qcaa_word_count.ingest import ingest
    from qcaa_word_count.interactive import resolve_all
    from qcaa_word_count.regions import (
        apply_regional_decisions, detect_regions,
    )
    from qcaa_word_count.tables import classify_table

    document = ingest(path)
    for table in document.tables:
        table.suggestion, table.suggestion_confidence = classify_table(table)
    detect_regions(document)
    apply_regional_decisions(document)
    _apply_span_decisions(document)
    detect_captions(document)
    document.footnotes = read_footnotes_from_docx(path)
    apply_footnote_suggestions(document)
    document.citations = collect_citations(document)
    apply_citation_exclusions(document)
    resolve_all(document, auto_accept=True)
    return count_words(document)


class TestBasicCounting(unittest.TestCase):
    def test_plain_paragraph(self):
        # "The quick brown fox jumps over the lazy dog" = 9 words
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["The quick brown fox jumps over the lazy dog"], path)
            result = _run(path)
            self.assertEqual(result["total"], 9)

    def test_numbers_excluded(self):
        # "There were 5 cats" -> "There were cats" = 3 words
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["There were 5 cats"], path)
            result = _run(path)
            self.assertEqual(result["total"], 3)

    def test_number_with_unit_excluded(self):
        # "The mass was 5kg" -> "The mass was" = 3 words
        # "5kg" is one token, excluded as a number-with-unit
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["The mass was 5kg"], path)
            result = _run(path)
            self.assertEqual(result["total"], 3)

    def test_unit_alone_counts(self):
        # "measured in kg" -> "measured in kg" = 3 words
        # "kg" is a unit and counts
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["measured in kg"], path)
            result = _run(path)
            self.assertEqual(result["total"], 3)

    def test_symbol_excluded(self):
        # "the + sign" -> "the sign" = 2 words
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["the + sign"], path)
            result = _run(path)
            self.assertEqual(result["total"], 2)

    def test_equation_excluded(self):
        # "x = 5" -> "x" is a variable in equation-heavy text, "=" is a symbol,
        # "5" is a number -> 0 words
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["x = 5"], path)
            result = _run(path)
            self.assertEqual(result["total"], 0)


class TestRegions(unittest.TestCase):
    def test_references_excluded(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx([
                "Introduction",
                "This is body text with five words.",
                "References",
                "Smith, J. (2020). A book. Publisher.",
            ], path)
            result = _run(path)
            # Body: "This is body text with five words" = 7
            # References: excluded
            self.assertEqual(result["total"], 7)


class TestCitations(unittest.TestCase):
    def test_apa_citation_excluded(self):
        # "(Smith, 2020)" is excluded, so the remaining count is the
        # text before and after it
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx([
                "Introduction",
                "Research shows this is true (Smith, 2020).",
            ], path)
            # "Research shows this is true" = 5
            result = _run(path)
            self.assertEqual(result["total"], 5)


class TestTables(unittest.TestCase):
    def test_data_table_excluded(self):
        # A table of numbers only -> raw data -> excluded
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            doc = DocxDocument()
            doc.add_paragraph("Introduction")
            doc.add_paragraph("Body text goes here.")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "1"
            table.cell(0, 1).text = "2"
            table.cell(1, 0).text = "3"
            table.cell(1, 1).text = "4"
            doc.save(path)
            result = _run(path)
            # "Body text goes here" = 4 words, table excluded
            self.assertEqual(result["total"], 4)


class TestCaptions(unittest.TestCase):
    def test_prose_figure_not_caption(self):
        # "Figure 3 shows the trend." is prose, not a caption
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx([
                "Introduction",
                "Figure 3 shows the trend.",
            ], path)
            result = _run(path)
            # "Figure shows the trend" = 4 (3 is a number, excluded)
            self.assertEqual(result["total"], 4)

    def test_labeled_caption_is_caption(self):
        # "Figure 3: Trend over time" is a caption and is excluded
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx([
                "Introduction",
                "Body text here.",
                "Figure 3: Trend over time",
            ], path)
            result = _run(path)
            # "Body text here" = 3, caption excluded
            self.assertEqual(result["total"], 3)


class TestDecisionLog(unittest.TestCase):
    def test_json_export_contains_all_blocks(self):
        from qcaa_word_count.output import export_decisions

        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.docx")
            _make_docx(["One two three."], path)
            result = _run(path)
            # Rebuild the document so we can export
            from qcaa_word_count.ingest import ingest
            document = ingest(path)
            json_out = export_decisions(document, result, format="json")
            self.assertIn("total", json_out)
            self.assertIn("blocks", json_out)


if __name__ == "__main__":
    unittest.main()