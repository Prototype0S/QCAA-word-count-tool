"""Ingest: .docx -> Document."""
from __future__ import annotations

import os

import docx
from docx.oxml.ns import qn

from .helpers import (
    cell_text,
    heading_level_from_style,
    make_spans_for_text,
    paragraph_is_list_item,
    paragraph_text,
    paragraph_text_boxes,
)
from .model import Block, BlockType, Document, Table


def ingest(path: str) -> Document:
    doc = docx.Document(path)
    document = Document(filename=os.path.basename(path))
    block_index = 0

    for para in doc.paragraphs:
        text, math_count = paragraph_text(para)
        if text.strip() or math_count:
            style_name = (para.style.name or "") if para.style else ""
            level = heading_level_from_style(style_name)
            if level:
                block_type = BlockType.HEADING
            elif paragraph_is_list_item(para):
                block_type = BlockType.LIST_ITEM
            else:
                block_type = BlockType.PARAGRAPH

            block = Block(
                block_type=block_type,
                block_index=block_index,
                style_name=style_name,
            )
            block.raw_text = text
            block.spans = make_spans_for_text(text, block_index=block_index)
            block.math_object_count = math_count
            document.blocks.append(block)
            block_index += 1

        for tb_text, tb_math in paragraph_text_boxes(para):
            if not tb_text.strip() and not tb_math:
                continue
            block = Block(
                block_type=BlockType.TEXT_BOX,
                block_index=block_index,
            )
            block.raw_text = tb_text
            block.spans = make_spans_for_text(tb_text, block_index=block_index)
            block.math_object_count = tb_math
            document.blocks.append(block)
            document.text_boxes.append(block)
            block_index += 1

    for t_idx, raw_table in enumerate(doc.tables):
        document.tables.append(_ingest_table(raw_table, t_idx))

    _count_drawings(doc, document)
    return document


def _ingest_table(raw_table, t_idx: int) -> Table:
    table = Table(table_index=t_idx)
    seen_tc: set[int] = set()

    for r_idx, raw_row in enumerate(raw_table.rows):
        row_blocks = []
        for c_idx, raw_cell in enumerate(raw_row.cells):
            tc_id = id(raw_cell._tc)
            synthetic_index = 100_000 + t_idx * 1000 + r_idx * 100 + c_idx

            if tc_id in seen_tc:
                # Merged cell — placeholder so row alignment is preserved
                cell_block = Block(
                    block_type=BlockType.TABLE_CELL,
                    block_index=synthetic_index,
                )
                cell_block.raw_text = ""
                cell_block.spans = []
                cell_block.math_object_count = 0
                row_blocks.append(cell_block)
                continue
            seen_tc.add(tc_id)

            text, math_count = cell_text(raw_cell)
            cell_block = Block(
                block_type=BlockType.TABLE_CELL,
                block_index=synthetic_index,
            )
            cell_block.raw_text = text
            cell_block.spans = make_spans_for_text(text, block_index=synthetic_index)
            cell_block.math_object_count = math_count
            row_blocks.append(cell_block)
        table.rows.append(row_blocks)

    table.header_row_index = _detect_header_row(raw_table)
    return table


def _detect_header_row(raw_table) -> int:
    """Return the row index of the header, from w:trPr/w:tblHeader.

    Falls back to 0 if no row is flagged as a header.
    """
    try:
        for r_idx, row in enumerate(raw_table.rows):
            trPr = row._tr.find(qn("w:trPr"))
            if trPr is not None:
                if trPr.find(qn("w:tblHeader")) is not None:
                    return r_idx
    except Exception:
        pass
    return 0


def _count_drawings(doc, document: Document) -> None:
    """Count drawing elements and those containing DrawingML text."""
    for node in doc.element.body.iter():
        if node.tag == qn("w:drawing"):
            document.drawings += 1
            for inner in node.iter():
                if inner.tag == qn("a:t") and (inner.text or "").strip():
                    document.drawings_with_text += 1
                    break